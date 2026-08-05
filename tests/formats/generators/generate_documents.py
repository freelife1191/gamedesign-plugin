#!/usr/bin/env python3

import json
import os
import tempfile
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def load_json(filename):
    return json.loads(Path(filename).read_text(encoding="utf-8"))


def content_for(case, source_root):
    result = load_json(source_root / "result.json")
    if case["product"] == "studio":
        economy = result["domain"]["economy"]
        experiment = result["domain"]["experiment"]
        rollback = result["domain"]["rollback"]
        return [
            {
                "title": "운영 판단 요약",
                "lead": f"일일 미션에서 소프트 재화를 공급하고 업그레이드에서 소비한다. 목표 보유량은 {economy['targetInventory']:,}이다.",
                "bullets": ["실물 가격 환산은 숨기지 않는다.", "확률은 공개한다.", "승인 범위 밖의 성과는 주장하지 않는다."],
                "sourcePointers": ["result.json#/domain/economy", "approval-snapshot.json#/responsibleGates/0"],
            },
            {
                "title": "경제 루프",
                "lead": "공급원 → 자원 → 소비처의 세 지점을 같은 주간 지표로 관찰한다.",
                "table": [["요소", "계약"], ["공급원", "일일 미션"], ["자원", "소프트 재화"], ["소비처", "업그레이드"], ["인플레이션 위험", "주간 공급-소비 차이"]],
                "sourcePointers": ["result.json#/domain/economy"],
            },
            {
                "title": "실험과 가드레일",
                "lead": "제한된 공급량 조정은 주간 인플레이션율을 높이지 않으면서 진행 체감을 개선한다.",
                "table": [["항목", "계약"], ["실험", experiment["experimentId"]], ["대조군", experiment["control"]], ["단일 변수", experiment["singleVariable"]], ["가드레일", experiment["guardrail"]], ["중단 조건", experiment["stopCondition"]]],
                "sourcePointers": ["result.json#/domain/experiment"],
            },
            {
                "title": "롤백 실행표",
                "lead": "가드레일 위반 시 자동으로 마지막 정상 설정을 복구한다.",
                "bullets": [f"트리거: {rollback['trigger']}", f"책임자: {rollback['owner']}", "복구 후 원장을 정합화하고 알린다."],
                "sourcePointers": ["result.json#/domain/rollback", "approval-snapshot.json#/responsibleGates/1"],
            },
        ]
    request = load_json(source_root / "request.json")
    roles = result["roleCandidates"]
    gaps = result["evidenceGaps"]
    brief = result["firstPortfolioBrief"]
    weeks = result["weeks"]
    return [
        {
            "title": "학습 판단 요약",
            "lead": f"12주 동안 주 {request['availableHoursPerWeek']}시간을 쓰며 작은 솔로 프로토타입으로 증거를 만든다.",
            "bullets": ["후보 역할은 시스템 디자인과 콘텐츠 디자인이다.", "현재 관심은 숙련 증거가 아니다.", "관찰 전 결과를 주장하지 않는다."],
            "sourcePointers": ["request.json#/availableHoursPerWeek"],
        },
        {
            "title": "역할 후보와 증거 공백",
            "lead": "역할 선택보다 먼저 각 역할이 요구하는 증거를 만든다.",
            "table": [["구분", "요구 증거"], [roles[0]["roleFamily"], roles[0]["proofArtifact"]], [roles[1]["roleFamily"], roles[1]["proofArtifact"]], [gaps[0]["gapId"], gaps[0]["proofArtifact"]], [gaps[1]["gapId"], gaps[1]["proofArtifact"]]],
            "sourcePointers": ["result.json#/roleCandidates", "result.json#/evidenceGaps"],
        },
        {
            "title": "12주 실행 리듬",
            "lead": "학습 → 연습 → 피드백을 매주 반복한다.",
            "table": [["주", "학습", "연습"]] + [[str(item["week"]), item["learning"], item["practice"]] for item in weeks],
            "sourcePointers": ["result.json#/weeks"],
        },
        {
            "title": "첫 포트폴리오 브리프",
            "lead": brief["problemUser"],
            "bullets": [brief["hypothesisIntent"], *brief["constraintsAlternatives"], brief["implementationTest"], "결과: 증거 대기. 세션 전 성과를 주장하지 않는다."],
            "sourcePointers": ["result.json#/firstPortfolioBrief"],
        },
    ]


def set_run_font(run, size=11, bold=False, color="0F172A"):
    run.font.name = "D2Coding"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "D2Coding")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "D2Coding")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "D2Coding")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hint"), "eastAsia")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def source_line(section):
    return " · ".join(section["sourcePointers"])


def table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def add_docx_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    widths = [2700, 6660] if len(rows[0]) == 2 else [900, 4230, 4230]
    for ri, values in enumerate(rows):
        for ci, value in enumerate(values):
            cell = table.cell(ri, ci)
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(value))
            set_run_font(run, size=9 if len(rows) > 8 else 10, bold=ri == 0)
            if ri == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E8EEF5")
                cell._tc.get_or_add_tcPr().append(shading)
    table_geometry(table, widths)


def make_docx(case, sections, output):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    styles = doc.styles
    for name, size, before, after, color in [("Normal", 11, 0, 6, "0F172A"), ("Title", 28, 0, 14, "0F172A"), ("Heading 1", 16, 18, 10, "2E74B5"), ("Heading 2", 13, 14, 7, "2E74B5")]:
        style = styles[name]
        style.font.name = "D2Coding"
        style_fonts = style._element.get_or_add_rPr().rFonts
        style_fonts.set(qn("w:ascii"), "D2Coding")
        style_fonts.set(qn("w:hAnsi"), "D2Coding")
        style_fonts.set(qn("w:eastAsia"), "D2Coding")
        style_fonts.set(qn("w:hint"), "eastAsia")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25
    header = sec.header.paragraphs[0]
    header.text = case["title"]
    set_run_font(header.runs[0], size=8, color="64748B")
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.add_run("Task 11 · format harness"), size=8, color="64748B")
    for index, section in enumerate(sections):
        if index:
            doc.add_page_break()
        if index == 0:
            for _ in range(4): doc.add_paragraph()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(p.add_run(case["product"].upper()), size=11, bold=True, color="3D8DFF")
            p = doc.add_paragraph()
            p.style = styles["Title"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title_size = 24 if case["product"] == "studio" else 28
            set_run_font(p.add_run(case["title"]), size=title_size, bold=True)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(p.add_run("근거 기반 실행 참고서"), size=13, color="475569")
            doc.add_paragraph()
        doc.add_heading(section["title"], level=1)
        p = doc.add_paragraph(section["lead"])
        for run in p.runs: set_run_font(run)
        for item in section.get("bullets", []):
            p = doc.add_paragraph(style="List Bullet")
            set_run_font(p.add_run(item))
        if section.get("table"):
            add_docx_table(doc, section["table"])
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        set_run_font(p.add_run(f"근거: {source_line(section)}"), size=8, color="64748B")
    props = doc.core_properties
    props.title = case["title"]
    props.subject = "Task 11 representative format output"
    props.author = "Game Design Plugin Suite"
    doc.save(output)
    patch_font_table(output)


def patch_font_table(filename):
    namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    with zipfile.ZipFile(filename, "r") as source:
        entries = {name: source.read(name) for name in source.namelist()}
    import xml.etree.ElementTree as ET
    root = ET.fromstring(entries["word/fontTable.xml"])
    if not any(font.get(f"{{{namespace}}}name") == "D2Coding" for font in root):
        font = ET.SubElement(root, f"{{{namespace}}}font", {f"{{{namespace}}}name": "D2Coding"})
        ET.SubElement(font, f"{{{namespace}}}altName", {f"{{{namespace}}}val": "Apple SD Gothic Neo"})
        ET.SubElement(font, f"{{{namespace}}}family", {f"{{{namespace}}}val": "modern"})
        ET.SubElement(font, f"{{{namespace}}}pitch", {f"{{{namespace}}}val": "fixed"})
        entries["word/fontTable.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    temporary = Path(tempfile.mkstemp(prefix="task11-docx-", suffix=".docx", dir=Path(filename).parent)[1])
    try:
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for name, data in entries.items(): target.writestr(name, data)
        os.replace(temporary, filename)
    finally:
        temporary.unlink(missing_ok=True)


def make_pdf(case, sections, font_path, output):
    pdfmetrics.registerFont(TTFont("D2Coding", font_path, subfontIndex=0))
    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleK", parent=styles["Title"], fontName="D2Coding", fontSize=25, leading=34, alignment=TA_CENTER, textColor=HexColor("#0F172A"), spaceAfter=18)
    h1 = ParagraphStyle("H1K", parent=styles["Heading1"], fontName="D2Coding", fontSize=18, leading=24, textColor=HexColor("#2E74B5"), spaceAfter=12)
    body = ParagraphStyle("BodyK", parent=styles["BodyText"], fontName="D2Coding", fontSize=10.5, leading=16, alignment=TA_LEFT, textColor=HexColor("#0F172A"), spaceAfter=8)
    small = ParagraphStyle("SmallK", parent=body, fontSize=8, leading=11, textColor=HexColor("#64748B"))
    def page(canvas, doc):
        canvas.saveState()
        canvas.setFont("D2Coding", 8)
        canvas.setFillColor(HexColor("#64748B"))
        canvas.drawString(inch, 0.55 * inch, case["title"])
        canvas.drawRightString(7.5 * inch, 0.55 * inch, str(doc.page))
        canvas.restoreState()
    story = []
    for index, section in enumerate(sections):
        if index: story.append(PageBreak())
        if index == 0:
            story += [Spacer(1, 1.25 * inch), Paragraph(case["product"].upper(), small), Paragraph(case["title"], title), Paragraph("근거 기반 실행 참고서", body), Spacer(1, 0.55 * inch)]
        story += [Paragraph(section["title"], h1), Paragraph(section["lead"], body)]
        for item in section.get("bullets", []): story.append(Paragraph(f"• {item}", body))
        if section.get("table"):
            rows = [[Paragraph(str(cell), small if len(section["table"]) > 8 else body) for cell in row] for row in section["table"]]
            widths = [1.65 * inch, 4.85 * inch] if len(rows[0]) == 2 else [0.6 * inch, 2.95 * inch, 2.95 * inch]
            table = Table(rows, colWidths=widths, repeatRows=1)
            table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), "D2Coding"), ("BACKGROUND", (0, 0), (-1, 0), HexColor("#E8EEF5")), ("GRID", (0, 0), (-1, -1), 0.5, HexColor("#94A3B8")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
            story.append(table)
        story += [Spacer(1, 10), Paragraph(f"근거: {source_line(section)}", small)]
    doc = SimpleDocTemplate(str(output), pagesize=letter, rightMargin=inch, leftMargin=inch, topMargin=0.8 * inch, bottomMargin=0.75 * inch, title=case["title"], author="Game Design Plugin Suite")
    doc.build(story, onFirstPage=page, onLaterPages=page)


def main():
    case_path, font_path, output_dir = map(Path, sys.argv[1:4])
    case = load_json(case_path)
    source_root = Path.cwd() / case["sourceRoot"]
    sections = content_for(case, source_root)
    output_dir.mkdir(parents=True, exist_ok=True)
    make_pdf(case, sections, str(font_path), output_dir / "brief.pdf")
    make_docx(case, sections, output_dir / "brief.docx")


if __name__ == "__main__":
    main()
